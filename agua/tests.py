import io
import tempfile
from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from PIL import Image
from rest_framework.test import APIClient

from .models import Socio, Medidor, Tarifa, Lectura, Cobro, Pago, CategoriaEgreso, Egreso
from .serializers import ReciboSocioSerializer

Usuario = get_user_model()


def imagen_png():
    buffer = io.BytesIO()
    Image.new('RGB', (40, 40), 'white').save(buffer, format='PNG')
    return SimpleUploadedFile('comprobante.png', buffer.getvalue(), content_type='image/png')


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class PagoEnRevisionTests(TestCase):
    """Pago por QR desde la app: En Revision -> aprobado / rechazado por el tesorero."""

    def setUp(self):
        Tarifa.objects.create(nombre='Base', costo_por_cubo=Decimal('2.00'),
                              cuota_fija=Decimal('30.00'), activa=True)
        self.usuario_socio = Usuario.objects.create_user(
            username='socio1', password='x', rol='socio')
        self.tesorero = Usuario.objects.create_user(
            username='teso', password='x', rol='tesorero')
        self.socio = Socio.objects.create(
            usuario=self.usuario_socio, ci='123456', nombre_completo='Juan Perez')
        medidor = Medidor.objects.create(socio=self.socio, numero_medidor='M-1')
        lectura = Lectura.objects.create(
            medidor=medidor,
            periodo=timezone.localdate().strftime('%Y-%m'),
            lectura_anterior=Decimal('0'), lectura_actual=Decimal('5'))
        self.cobro = Cobro.objects.get(lectura=lectura)
        self.api = APIClient()
        self.api.force_authenticate(self.usuario_socio)
        self.url_ocr = f'/api/mi-cuenta/{self.cobro.pk}/validar-pago-ocr/'

    def enviar_comprobante(self, texto='PAGO QR 30.00 BS NRO 987654321'):
        with patch('agua.views.llamar_ocr_space',
                   return_value={'exitoso': True, 'texto': texto}):
            return self.api.post(self.url_ocr, {'comprobante': imagen_png()}, format='multipart')

    def revisar(self, pago, **datos):
        self.client.force_login(self.tesorero)
        return self.client.post(reverse('pago_revisar', args=[pago.pk]), datos)

    def test_comprobante_valido_queda_en_revision(self):
        self.assertEqual(self.cobro.monto_total, Decimal('30.00'))
        respuesta = self.enviar_comprobante()

        self.assertEqual(respuesta.status_code, 200, respuesta.data)
        self.assertEqual(respuesta.data['estado'], 'En Revision')
        self.cobro.refresh_from_db()
        self.assertEqual(self.cobro.estado_pago, 'En Revision')
        pago = self.cobro.pagos.get()
        self.assertEqual(pago.estado, 'En Revision')
        self.assertEqual(pago.nro_transaccion, '987654321')
        # Todavía no cuenta como cobrado
        self.assertEqual(self.cobro.total_pagado_aprobado(), Decimal('0.00'))

    def test_no_acepta_segundo_comprobante_mientras_esta_en_revision(self):
        self.enviar_comprobante()
        respuesta = self.enviar_comprobante()

        self.assertEqual(respuesta.status_code, 400)
        self.assertEqual(respuesta.data['codigo'], 'YA_EN_REVISION')
        self.assertEqual(self.cobro.pagos.count(), 1)

    def test_archivo_que_no_es_imagen_pide_reenvio(self):
        archivo = SimpleUploadedFile('nota.txt', b'hola', content_type='text/plain')
        respuesta = self.api.post(self.url_ocr, {'comprobante': archivo}, format='multipart')

        self.assertEqual(respuesta.status_code, 400)
        self.assertEqual(respuesta.data['codigo'], 'REENVIAR_COMPROBANTE')
        self.assertFalse(self.cobro.pagos.exists())

    def test_sin_archivo_pide_reenvio(self):
        respuesta = self.api.post(self.url_ocr, {}, format='multipart')

        self.assertEqual(respuesta.status_code, 400)
        self.assertEqual(respuesta.data['codigo'], 'REENVIAR_COMPROBANTE')

    def test_imagen_sin_el_monto_pide_reenvio(self):
        respuesta = self.enviar_comprobante(texto='UNA FOTO CUALQUIERA')

        self.assertEqual(respuesta.status_code, 400)
        self.assertEqual(respuesta.data['codigo'], 'REENVIAR_COMPROBANTE')
        self.cobro.refresh_from_db()
        self.assertEqual(self.cobro.estado_pago, 'Pendiente')

    def test_aprobar_cancela_el_cobro(self):
        self.enviar_comprobante()
        pago = self.cobro.pagos.get()

        self.revisar(pago, accion='aprobar')

        pago.refresh_from_db()
        self.cobro.refresh_from_db()
        self.assertEqual(pago.estado, 'Aprobado')
        self.assertEqual(pago.revisado_por, self.tesorero)
        self.assertEqual(self.cobro.estado_pago, 'Cancelado')
        datos = ReciboSocioSerializer(self.cobro).data
        self.assertEqual(datos['metodo_pago'], 'qr')

    def test_rechazar_devuelve_el_cobro_a_pendiente_y_avisa_al_socio(self):
        self.enviar_comprobante()
        pago = self.cobro.pagos.get()

        self.revisar(pago, accion='rechazar', motivo_rechazo='Foto borrosa')

        pago.refresh_from_db()
        self.cobro.refresh_from_db()
        self.assertEqual(pago.estado, 'Rechazado')
        self.assertEqual(self.cobro.estado_pago, 'Pendiente')
        datos = ReciboSocioSerializer(self.cobro).data
        self.assertTrue(datos['comprobante_rechazado'])
        self.assertEqual(datos['motivo_rechazo'], 'Foto borrosa')
        # La app usa este dato para avisar una sola vez por cada rechazo
        self.assertEqual(datos['rechazo_id'], str(pago.pk))
        # El socio puede volver a enviar otro comprobante
        self.assertEqual(self.enviar_comprobante().status_code, 200)

    def test_rechazar_exige_motivo(self):
        self.enviar_comprobante()
        pago = self.cobro.pagos.get()

        self.revisar(pago, accion='rechazar', motivo_rechazo='')

        pago.refresh_from_db()
        self.assertEqual(pago.estado, 'En Revision')

    def test_socio_no_puede_revisar(self):
        self.enviar_comprobante()
        pago = self.cobro.pagos.get()
        self.client.force_login(self.usuario_socio)

        self.client.post(reverse('pago_revisar', args=[pago.pk]), {'accion': 'aprobar'})

        pago.refresh_from_db()
        self.assertEqual(pago.estado, 'En Revision')

    def test_pago_en_efectivo_cancela_directo(self):
        Pago.objects.create(recibo=self.cobro, monto_pagado=Decimal('30.00'),
                            metodo_pago='efectivo', registrado_por=self.tesorero)

        self.cobro.refresh_from_db()
        self.assertEqual(self.cobro.estado_pago, 'Cancelado')
        datos = ReciboSocioSerializer(self.cobro).data
        self.assertEqual(datos['metodo_pago'], 'efectivo')
        self.assertEqual(datos['mensaje_estado'], 'Cancelado - pago en efectivo')

    def test_no_hay_recargo_mientras_esta_en_revision(self):
        self.enviar_comprobante()
        self.cobro.refresh_from_db()
        dentro_de_un_anio = timezone.localdate().replace(year=timezone.localdate().year + 1)

        self.assertFalse(self.cobro.aplicar_recargo_automatico(hoy=dentro_de_un_anio))
        self.assertEqual(self.cobro.recargo_falta_pago, Decimal('0.00'))

    def test_pantallas_de_revision_cargan(self):
        self.enviar_comprobante()
        pago = self.cobro.pagos.get()
        self.client.force_login(self.tesorero)

        for url in (reverse('pagos_revision'), reverse('pago_revisar', args=[pago.pk]),
                    reverse('pagos_lista'), reverse('cobro_detalle', args=[self.cobro.pk])):
            self.assertEqual(self.client.get(url).status_code, 200, url)


class TelefonoSocioTests(TestCase):
    """Teléfono del socio: Bolivia sin código, otros países con su código."""

    def setUp(self):
        admin = Usuario.objects.create_user(username='adm', password='x', rol='admin')
        self.client.force_login(admin)

    def crear(self, **datos):
        base = {'ci': '555', 'nombre_completo': 'Ana Lopez'}
        base.update(datos)
        self.client.post(reverse('socio_crear'), base)
        return Socio.objects.filter(ci='555').first()

    def test_bolivia_se_guarda_sin_codigo(self):
        socio = self.crear(codigo_pais='591', telefono='71234567')
        self.assertEqual(socio.telefono, '71234567')

    def test_otro_pais_se_guarda_con_codigo(self):
        socio = self.crear(codigo_pais='54', telefono='91123456789')
        self.assertEqual(socio.telefono, '+54 91123456789')

    def test_codigo_escrito_a_mano(self):
        socio = self.crear(codigo_pais='otro', codigo_pais_otro='49', telefono='15112345678')
        self.assertEqual(socio.telefono, '+49 15112345678')

    def test_bolivia_con_mas_de_8_digitos_se_rechaza(self):
        self.assertIsNone(self.crear(codigo_pais='591', telefono='712345678'))

    def test_sin_telefono(self):
        self.assertIsNone(self.crear(codigo_pais='591', telefono='').telefono)

    def test_editar_muestra_codigo_y_numero_separados(self):
        socio = self.crear(codigo_pais='54', telefono='91123456789')
        respuesta = self.client.get(reverse('socio_editar', args=[socio.pk]))
        self.assertEqual(respuesta.context['tel_codigo'], '54')
        self.assertEqual(respuesta.context['tel_numero'], '91123456789')


class LecturaOdometroTests(TestCase):
    """Interpretación del odómetro: 7 dígitos, los 2 últimos son decimales."""

    SERIE = 'A18S801878'

    def leer(self, texto, anterior=480, techo=530):
        from .views_web import _extraer_posible_lectura
        return _extraer_posible_lectura(texto, self.SERIE, Decimal(anterior), techo)

    def test_odometro_completo_de_7_digitos(self):
        self.assertEqual(self.leer('Itron\nA18S801878\n0048327\nm3'), 483)

    def test_odometro_partido_en_negros_y_rojos(self):
        self.assertEqual(self.leer('A18S801878\n00483 27'), 483)

    def test_odometro_con_digitos_separados(self):
        self.assertEqual(self.leer('0 0 4 8 3 2 7\nA18S801878'), 483)

    def test_solo_los_5_digitos_negros(self):
        self.assertEqual(self.leer('A18S801878\n00483'), 483)

    def test_letras_parecidas_a_digitos(self):
        self.assertEqual(self.leer('A18S801878\nOO48327'), 483)

    def test_no_usa_los_digitos_del_numero_de_serie(self):
        self.assertIsNone(self.leer('A18S801878\nIndustria Brasileira', anterior=0, techo=50))

    def test_palabras_no_se_toman_como_numeros(self):
        self.assertIsNone(self.leer('ISO 4064 B\nMULTIMAG CYBLE', anterior=100, techo=160))

    def test_fuera_de_rango_se_descarta(self):
        self.assertIsNone(self.leer('0098327'))

    def test_primera_lectura_acepta_el_odometro_sin_rango(self):
        self.assertEqual(self.leer('A18S801878\n0048327', anterior=0, techo=None), 483)

    def test_primera_lectura_no_adivina_con_numeros_sueltos(self):
        self.assertIsNone(self.leer('A18S801878\n1.5 m3/h\n5 4 3', anterior=0, techo=None))


class DetalleSocioTests(TestCase):
    """La ficha del socio debe abrir aunque el socio ya tenga cobros."""

    def test_detalle_de_socio_con_cobros_carga(self):
        Tarifa.objects.create(nombre='Base', costo_por_cubo=Decimal('2.00'),
                              cuota_fija=Decimal('30.00'), activa=True)
        admin = Usuario.objects.create_user(username='adm2', password='x', rol='admin')
        socio = Socio.objects.create(ci='777', nombre_completo='Luis Rojas')
        medidor = Medidor.objects.create(socio=socio, numero_medidor='M-7')
        Lectura.objects.create(medidor=medidor, periodo=timezone.localdate().strftime('%Y-%m'),
                               lectura_anterior=Decimal('0'), lectura_actual=Decimal('5'))
        self.client.force_login(admin)

        respuesta = self.client.get(reverse('socio_detalle', args=[socio.pk]))

        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, reverse('cobro_imprimir_termico', args=[socio.recibos.first().pk]))


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class DatosDePruebaTests(TestCase):
    """Comandos crear_datos_prueba / eliminar_datos_prueba."""

    def setUp(self):
        Tarifa.objects.create(nombre='Base', costo_por_cubo=Decimal('2.00'),
                              cuota_fija=Decimal('30.00'), multa_atraso=Decimal('5.00'),
                              activa=True)

    def test_crea_un_recibo_en_cada_estado_y_luego_los_borra(self):
        from django.core.management import call_command
        from .management.commands.crear_datos_prueba import CI_PRUEBA

        call_command('crear_datos_prueba', stdout=io.StringIO())

        socio = Socio.objects.get(ci=CI_PRUEBA)
        estados = list(
            socio.recibos.order_by('lectura__periodo').values_list('estado_pago', flat=True)
        )
        self.assertEqual(
            estados,
            ['Cancelado', 'Cancelado', 'Vencido', 'Vencido', 'Pendiente', 'En Revision'],
        )
        vencidos = socio.recibos.filter(estado_pago='Vencido').order_by('lectura__periodo')
        self.assertEqual(vencidos[0].recargo_falta_pago, Decimal('10.00'))  # 2 meses
        self.assertEqual(vencidos[1].recargo_falta_pago, Decimal('5.00'))   # 1 mes
        self.assertTrue(ReciboSocioSerializer(vencidos[1]).data['comprobante_rechazado'])
        self.assertTrue(Usuario.objects.get(username=CI_PRUEBA).check_password(CI_PRUEBA))

        call_command('eliminar_datos_prueba', stdout=io.StringIO())

        self.assertFalse(Socio.objects.filter(ci=CI_PRUEBA).exists())
        self.assertFalse(Usuario.objects.filter(username=CI_PRUEBA).exists())
        self.assertFalse(Cobro.objects.exists())
        self.assertFalse(Pago.objects.exists())

    def test_no_borra_a_un_socio_real_con_el_mismo_ci(self):
        from django.core.management import call_command
        from django.core.management.base import CommandError
        from .management.commands.crear_datos_prueba import CI_PRUEBA

        Socio.objects.create(ci=CI_PRUEBA, nombre_completo='Persona Real')

        with self.assertRaises(CommandError):
            call_command('eliminar_datos_prueba', stdout=io.StringIO())
        self.assertTrue(Socio.objects.filter(ci=CI_PRUEBA).exists())


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class EgresosTests(TestCase):
    """Gastos del comité: registro, anulación, permisos y balance."""

    def setUp(self):
        self.admin = Usuario.objects.create_user(username='admin1', password='x', rol='admin')
        self.tesorero = Usuario.objects.create_user(username='teso', password='x', rol='tesorero')
        self.lector = Usuario.objects.create_user(username='lector1', password='x', rol='lector')
        # Las categorías iniciales las crea la migración.
        self.categoria = CategoriaEgreso.objects.get(nombre='Limpieza')
        self.client.force_login(self.tesorero)

    def datos(self, **cambios):
        datos = {
            'fecha': timezone.localdate().isoformat(),
            'categoria': self.categoria.pk,
            'descripcion': 'Limpieza del tanque',
            'pagado_a': 'Pedro Mamani',
            'monto': '150.50',
        }
        datos.update(cambios)
        return datos

    def crear_egreso(self, monto='100.00', **campos):
        return Egreso.objects.create(
            categoria=self.categoria, descripcion='Gasto', monto=Decimal(monto), **campos)

    def test_tesorero_registra_un_egreso_con_foto(self):
        respuesta = self.client.post(
            reverse('egreso_crear'), self.datos(foto_respaldo=imagen_png()))

        self.assertRedirects(respuesta, reverse('egresos_lista'))
        egreso = Egreso.objects.get()
        self.assertEqual(egreso.monto, Decimal('150.50'))
        self.assertEqual(egreso.registrado_por, self.tesorero)
        self.assertTrue(egreso.foto_respaldo)

    def test_no_guarda_monto_invalido_ni_fecha_futura(self):
        manana = (timezone.localdate() + timezone.timedelta(days=1)).isoformat()
        for cambios in ({'monto': '0'}, {'monto': 'abc'}, {'fecha': manana}, {'categoria': ''}):
            respuesta = self.client.post(reverse('egreso_crear'), self.datos(**cambios))
            self.assertEqual(respuesta.status_code, 200)
        self.assertFalse(Egreso.objects.exists())

    def test_anular_exige_motivo_y_deja_de_sumar(self):
        egreso = self.crear_egreso('100.00')
        self.crear_egreso('40.00')
        url = reverse('egreso_anular', args=[egreso.pk])

        self.client.post(url, {'motivo_anulacion': ''})
        egreso.refresh_from_db()
        self.assertFalse(egreso.anulado)

        self.client.post(url, {'motivo_anulacion': 'Registrado dos veces'})
        egreso.refresh_from_db()
        self.assertTrue(egreso.anulado)

        respuesta = self.client.get(reverse('egresos_lista'))
        self.assertEqual(respuesta.context['total'], Decimal('40.00'))
        self.assertEqual(respuesta.context['cantidad'], 1)

    def test_egreso_anulado_no_se_puede_editar(self):
        egreso = self.crear_egreso(anulado=True, motivo_anulacion='Error')
        self.client.post(reverse('egreso_editar', args=[egreso.pk]), self.datos(monto='999'))
        egreso.refresh_from_db()
        self.assertEqual(egreso.monto, Decimal('100.00'))

    def test_la_lista_no_falla_si_se_borro_el_usuario_que_registro(self):
        self.crear_egreso(registrado_por=None)
        self.assertEqual(self.client.get(reverse('egresos_lista')).status_code, 200)
        self.assertEqual(self.client.get(reverse('egresos_excel')).status_code, 200)

    def test_el_lector_no_entra_a_egresos(self):
        self.client.force_login(self.lector)
        respuesta = self.client.post(reverse('egreso_crear'), self.datos())
        self.assertEqual(respuesta.status_code, 302)
        self.assertFalse(Egreso.objects.exists())

    def test_solo_el_admin_administra_categorias(self):
        url = reverse('categorias_egreso')
        self.client.post(url, {'accion': 'crear', 'nombre': 'Compra de cloro'})
        self.assertFalse(CategoriaEgreso.objects.filter(nombre='Compra de cloro').exists())

        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(url).status_code, 200)
        self.client.post(url, {'accion': 'crear', 'nombre': 'Compra de cloro'})
        self.client.post(url, {'accion': 'crear', 'nombre': 'compra de CLORO'})  # repetida
        self.assertEqual(CategoriaEgreso.objects.filter(nombre__iexact='compra de cloro').count(), 1)

        self.client.post(url, {'accion': 'cambiar_estado', 'pk': self.categoria.pk})
        self.categoria.refresh_from_db()
        self.assertFalse(self.categoria.activo)
        # Una categoría inactiva ya no se puede usar en egresos nuevos... pero
        # sigue apareciendo en el filtro de la lista para ver su historial.
        respuesta = self.client.get(reverse('egreso_crear'))
        self.assertNotIn(self.categoria, respuesta.context['categorias'])

    def test_balance_resta_los_egresos_a_lo_recaudado(self):
        Tarifa.objects.create(nombre='Base', costo_por_cubo=Decimal('2.00'),
                              cuota_fija=Decimal('30.00'), activa=True)
        socio = Socio.objects.create(ci='123456', nombre_completo='Juan Perez')
        medidor = Medidor.objects.create(socio=socio, numero_medidor='M-1')
        lectura = Lectura.objects.create(
            medidor=medidor, periodo=timezone.localdate().strftime('%Y-%m'),
            lectura_anterior=Decimal('0'), lectura_actual=Decimal('5'))
        Pago.objects.create(recibo=Cobro.objects.get(lectura=lectura),
                            monto_pagado=Decimal('30.00'), metodo_pago='efectivo')
        self.crear_egreso('10.00')
        self.crear_egreso('500.00', anulado=True, motivo_anulacion='Error')

        respuesta = self.client.get(reverse('reporte_balance'))

        totales = respuesta.context['totales']
        self.assertEqual(totales['ingresos'], Decimal('30.00'))
        self.assertEqual(totales['egresos'], Decimal('10.00'))
        self.assertEqual(totales['saldo'], Decimal('20.00'))
        mes = respuesta.context['meses'][timezone.localdate().month - 1]
        self.assertEqual(mes['saldo'], Decimal('20.00'))
        self.assertEqual(self.client.get(reverse('reporte_balance_excel')).status_code, 200)

        # El egreso también se resta en el panel y en los demás reportes.
        panel = self.client.get(reverse('dashboard'))
        self.assertEqual(panel.context['ingresos_acumulados'], Decimal('30.00'))
        self.assertEqual(panel.context['egresos_acumulados'], Decimal('10.00'))
        self.assertEqual(panel.context['saldo_caja'], Decimal('20.00'))
        self.assertEqual(panel.context['saldo_mes'], Decimal('20.00'))
        for nombre, clave in (('reporte_recaudacion', 'saldo'),
                              ('reporte_mensual', 'saldo'),
                              ('reporte_anual', 'saldo_anual')):
            reporte = self.client.get(reverse(nombre))
            self.assertEqual(reporte.status_code, 200)
            self.assertEqual(reporte.context[clave], Decimal('20.00'))

        # Y el reporte general lo explica con frases.
        general = self.client.get(reverse('reporte_general'))
        self.assertEqual(general.status_code, 200)
        self.assertEqual(general.context['totales']['saldo'], Decimal('20.00'))
        titulos = [frase['titulo'] for frase in general.context['interpretacion']]
        self.assertIn('Saldo a favor', titulos)
        self.assertIn('En qué se gasta', titulos)

    def test_reporte_general_sin_datos_no_falla(self):
        respuesta = self.client.get(reverse('reporte_general'), {'anio': '2001'})
        self.assertEqual(respuesta.status_code, 200)
        titulos = [frase['titulo'] for frase in respuesta.context['interpretacion']]
        self.assertIn('Sin movimiento', titulos)

    def test_el_backup_incluye_los_egresos(self):
        import openpyxl
        self.crear_egreso('75.00', pagado_a='Ferretería')

        self.assertEqual(self.client.get(reverse('backup_vista')).status_code, 200)
        respuesta = self.client.post(reverse('backup_excel'), {'password_confirmacion': 'x'})

        self.assertEqual(respuesta.status_code, 200)
        libro = openpyxl.load_workbook(io.BytesIO(respuesta.content))
        self.assertIn('Egresos', libro.sheetnames)
        self.assertIn('Categorias Egreso', libro.sheetnames)
        fila = [celda.value for celda in libro['Egresos'][2]]
        self.assertIn('Ferretería', fila)
        self.assertIn(75.0, fila)
        resumen = {fila[0].value: fila[1].value for fila in libro['Resumen'].iter_rows(min_row=7)}
        self.assertEqual(resumen['Total Egresos (Bs)'], 75.0)

    def test_importar_egresos_desde_la_plantilla_sin_duplicar(self):
        import openpyxl
        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.title = 'Egresos'
        hoja.append(['fecha', 'categoria', 'descripcion', 'pagado_a', 'monto', 'nro_factura_recibo'])
        hoja.append(['2026-03-05', 'Limpieza', 'Limpieza del tanque', 'Pedro', 120, 'R-15'])
        hoja.append(['2026-03-06', 'Compra de cloro', 'Cloro granulado', '', 80.5, ''])
        contenido = io.BytesIO()
        libro.save(contenido)

        for _ in range(2):  # subir dos veces el mismo archivo no repite los gastos
            archivo = SimpleUploadedFile('egresos.xlsx', contenido.getvalue())
            self.client.post(reverse('importar_excel'),
                             {'password_confirmacion': 'x', 'archivo_excel': archivo})

        self.assertEqual(Egreso.objects.count(), 2)
        self.assertTrue(CategoriaEgreso.objects.filter(nombre='Compra de cloro').exists())
        self.assertEqual(Egreso.objects.get(nro_respaldo='R-15').monto, Decimal('120.00'))

    def test_pagos_lista_no_falla_con_un_pago_sin_usuario(self):
        """Error visto en producción: un pago con registrado_por vacío tumbaba /pagos/."""
        Tarifa.objects.create(nombre='Base', costo_por_cubo=Decimal('2.00'),
                              cuota_fija=Decimal('30.00'), activa=True)
        socio = Socio.objects.create(ci='123456', nombre_completo='Juan Perez',
                                     codigo_cliente='SOC-777')
        medidor = Medidor.objects.create(socio=socio, numero_medidor='M-1')
        lectura = Lectura.objects.create(
            medidor=medidor, periodo=timezone.localdate().strftime('%Y-%m'),
            lectura_anterior=Decimal('0'), lectura_actual=Decimal('5'))
        cobro = Cobro.objects.get(lectura=lectura)
        pago = Pago.objects.create(recibo=cobro, monto_pagado=Decimal('30.00'),
                                   registrado_por=None, estado='En Revision')

        # Además, en cobros y en la revisión de pagos se ve el código de socio.
        for url in (reverse('pagos_lista'),
                    reverse('cobro_detalle', args=[cobro.pk]),
                    reverse('cobros_lista'),
                    reverse('pagos_revision'),
                    reverse('pago_revisar', args=[pago.pk])):
            respuesta = self.client.get(url)
            self.assertEqual(respuesta.status_code, 200, url)
            self.assertContains(respuesta, 'SOC-777')
