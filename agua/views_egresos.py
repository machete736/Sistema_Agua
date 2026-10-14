import io
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Q, Sum
from django.db.models.functions import ExtractMonth
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from django.utils.dateparse import parse_date

from .models import Afiliacion, CategoriaEgreso, Egreso, Pago
from .views_web import es_admin_o_tesorero, es_solo_admin, _metadata_reporte, MESES_NOMBRE


CERO = Decimal('0.00')


# =============================================================
# EGRESOS — lista, registro, edición y anulación
# =============================================================

def _egresos_filtrados(request):
    """Egresos según los filtros de la pantalla (los usa la lista y el Excel)."""
    filtros = {
        'fecha_inicio': request.GET.get('fecha_inicio', '').strip(),
        'fecha_fin': request.GET.get('fecha_fin', '').strip(),
        'categoria': request.GET.get('categoria', '').strip(),
        'q': request.GET.get('q', '').strip(),
    }

    egresos = Egreso.objects.select_related('categoria', 'registrado_por')

    fi = parse_date(filtros['fecha_inicio']) if filtros['fecha_inicio'] else None
    if fi:
        egresos = egresos.filter(fecha__gte=fi)

    ff = parse_date(filtros['fecha_fin']) if filtros['fecha_fin'] else None
    if ff:
        egresos = egresos.filter(fecha__lte=ff)

    if filtros['categoria'].isdigit():
        egresos = egresos.filter(categoria_id=int(filtros['categoria']))

    if filtros['q']:
        egresos = egresos.filter(
            Q(descripcion__icontains=filtros['q']) |
            Q(pagado_a__icontains=filtros['q']) |
            Q(nro_respaldo__icontains=filtros['q'])
        )

    return egresos, filtros


def _total_egresos(egresos):
    """Suma solo los egresos vigentes: los anulados no cuentan."""
    return egresos.filter(anulado=False).aggregate(total=Sum('monto'))['total'] or CERO


@login_required
@es_admin_o_tesorero
def egresos_lista(request):
    egresos, filtros = _egresos_filtrados(request)

    pagina = Paginator(egresos, 50).get_page(request.GET.get('page'))

    # Para que "Anterior / Siguiente" y el Excel conserven los filtros.
    parametros = request.GET.copy()
    parametros.pop('page', None)

    return render(request, 'egresos/lista.html', {
        'pagina': pagina,
        'total': _total_egresos(egresos),
        'cantidad': egresos.filter(anulado=False).count(),
        'categorias': CategoriaEgreso.objects.all(),
        'parametros': parametros.urlencode(),
        **filtros,
    })


def _guardar_egreso(request, egreso):
    """
    Valida el formulario y guarda el egreso. Devuelve True si se guardó;
    si no, deja el mensaje de error y devuelve False.
    """
    fecha = parse_date(request.POST.get('fecha', '').strip() or '')
    categoria_id = request.POST.get('categoria', '').strip()
    descripcion = request.POST.get('descripcion', '').strip()
    pagado_a = request.POST.get('pagado_a', '').strip()
    nro_respaldo = request.POST.get('nro_respaldo', '').strip()
    foto = request.FILES.get('foto_respaldo')

    try:
        monto = Decimal(request.POST.get('monto', '').strip().replace(',', '.'))
    except InvalidOperation:
        monto = None

    categoria = None
    if categoria_id.isdigit():
        categoria = CategoriaEgreso.objects.filter(pk=int(categoria_id)).first()

    if not fecha:
        messages.error(request, 'La fecha del gasto es obligatoria.')
    elif fecha > timezone.localdate():
        messages.error(request, 'La fecha del gasto no puede ser una fecha futura.')
    elif not categoria:
        messages.error(request, 'Elige la categoría del egreso.')
    elif not descripcion:
        messages.error(request, 'Escribe en qué se gastó.')
    elif monto is None or monto <= 0:
        messages.error(request, 'El monto debe ser un número mayor a cero.')
    elif monto >= Decimal('100000000'):
        messages.error(request, 'El monto es demasiado grande.')
    elif foto and not (foto.content_type or '').startswith('image/'):
        messages.error(request, 'El respaldo debe ser una imagen (foto de la factura o recibo).')
    else:
        egreso.fecha = fecha
        egreso.categoria = categoria
        egreso.descripcion = descripcion[:255]
        egreso.pagado_a = pagado_a[:150]
        egreso.nro_respaldo = nro_respaldo[:50]
        egreso.monto = monto.quantize(Decimal('0.01'))
        if foto:  # Solo se cambia la foto si se subió una nueva
            egreso.foto_respaldo = foto
        if not egreso.pk:
            egreso.registrado_por = request.user
        egreso.save()
        return True

    return False


def _contexto_form_egreso(request, accion, egreso=None):
    """
    Datos del formulario. Si el guardado falló, se vuelve a mostrar lo que
    la persona escribió para que no tenga que llenar todo de nuevo.
    """
    datos = request.POST if request.method == 'POST' else None
    categorias = CategoriaEgreso.objects.filter(
        Q(activo=True) | Q(pk=egreso.categoria_id if egreso else None)
    )
    return {
        'accion': accion,
        'egreso': egreso,
        'categorias': categorias,
        'valor_fecha': datos.get('fecha', '') if datos else (
            egreso.fecha.isoformat() if egreso else timezone.localdate().isoformat()),
        'valor_categoria': datos.get('categoria', '') if datos else (
            str(egreso.categoria_id) if egreso else ''),
        'valor_descripcion': datos.get('descripcion', '') if datos else (
            egreso.descripcion if egreso else ''),
        'valor_pagado_a': datos.get('pagado_a', '') if datos else (
            egreso.pagado_a if egreso else ''),
        'valor_nro_respaldo': datos.get('nro_respaldo', '') if datos else (
            egreso.nro_respaldo if egreso else ''),
        'valor_monto': datos.get('monto', '') if datos else (
            str(egreso.monto) if egreso else ''),
    }


@login_required
@es_admin_o_tesorero
def egreso_crear(request):
    if request.method == 'POST':
        egreso = Egreso()
        if _guardar_egreso(request, egreso):
            messages.success(request, f'Egreso de Bs {egreso.monto} registrado correctamente.')
            return redirect('egresos_lista')

    return render(request, 'egresos/form.html', _contexto_form_egreso(request, 'Registrar'))


@login_required
@es_admin_o_tesorero
def egreso_editar(request, pk):
    egreso = get_object_or_404(Egreso, pk=pk)

    if egreso.anulado:
        messages.warning(request, 'Este egreso está anulado y ya no se puede modificar.')
        return redirect('egresos_lista')

    if request.method == 'POST':
        if _guardar_egreso(request, egreso):
            messages.success(request, 'Egreso actualizado correctamente.')
            return redirect('egresos_lista')

    return render(request, 'egresos/form.html', _contexto_form_egreso(request, 'Editar', egreso))


@login_required
@es_admin_o_tesorero
def egreso_anular(request, pk):
    egreso = get_object_or_404(Egreso.objects.select_related('categoria'), pk=pk)

    if egreso.anulado:
        messages.warning(request, 'Este egreso ya estaba anulado.')
        return redirect('egresos_lista')

    if request.method == 'POST':
        motivo = request.POST.get('motivo_anulacion', '').strip()
        if not motivo:
            messages.error(request, 'Escribe el motivo de la anulación.')
        else:
            egreso.anulado = True
            egreso.motivo_anulacion = motivo[:255]
            egreso.save()
            messages.success(request, f'Egreso N.º {egreso.pk} anulado. Ya no suma en los totales.')
            return redirect('egresos_lista')

    return render(request, 'egresos/anular.html', {'egreso': egreso})


# =============================================================
# CATEGORÍAS DE EGRESO — solo el administrador
# =============================================================

@login_required
@es_solo_admin
def categorias_egreso(request):
    if request.method == 'POST':
        accion = request.POST.get('accion')
        nombre = ' '.join(request.POST.get('nombre', '').split())[:80]

        if accion == 'crear':
            if not nombre:
                messages.error(request, 'Escribe el nombre de la categoría.')
            elif CategoriaEgreso.objects.filter(nombre__iexact=nombre).exists():
                messages.error(request, f'Ya existe la categoría "{nombre}".')
            else:
                CategoriaEgreso.objects.create(nombre=nombre)
                messages.success(request, f'Categoría "{nombre}" creada.')

        elif accion in ('renombrar', 'cambiar_estado'):
            categoria = get_object_or_404(CategoriaEgreso, pk=request.POST.get('pk') or 0)
            if accion == 'cambiar_estado':
                categoria.activo = not categoria.activo
                categoria.save()
                estado = 'activada' if categoria.activo else 'desactivada'
                messages.success(request, f'Categoría "{categoria.nombre}" {estado}.')
            elif not nombre:
                messages.error(request, 'El nombre de la categoría no puede quedar vacío.')
            elif CategoriaEgreso.objects.filter(nombre__iexact=nombre).exclude(pk=categoria.pk).exists():
                messages.error(request, f'Ya existe otra categoría llamada "{nombre}".')
            else:
                categoria.nombre = nombre
                categoria.save()
                messages.success(request, 'Categoría actualizada.')

        return redirect('categorias_egreso')

    return render(request, 'egresos/categorias.html', {
        'categorias': CategoriaEgreso.objects.all(),
    })


# =============================================================
# EXCEL
# =============================================================

_AZUL = PatternFill("solid", fgColor="0D4F7C")
_GRIS = PatternFill("solid", fgColor="F2F2F2")
_BORDE = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
_FORMATO_BS = '#,##0.00'


def _hoja_con_tabla(ws, titulo, columnas, filas, fila_total=None, columnas_monto=()):
    """
    Llena una hoja con: nombre del comité, título, encabezados, filas y
    (opcional) una fila de totales. `columnas_monto` son las posiciones
    (desde 1) que llevan formato de dinero.
    """
    ws.append(['COMITÉ DE AGUA POTABLE JUNTA VECINAL "PARAÍSO"'])
    ws.append([titulo])
    ws.append([])
    ws['A1'].font = Font(bold=True, size=12)
    ws['A2'].font = Font(bold=True, size=11)

    ws.append(columnas)
    fila_encabezado = ws.max_row
    for cell in ws[fila_encabezado]:
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = _AZUL
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = _BORDE

    for fila in filas:
        ws.append(fila)
        for cell in ws[ws.max_row]:
            cell.border = _BORDE
            if ws.max_row % 2 == 0:
                cell.fill = _GRIS

    if fila_total:
        ws.append(fila_total)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.border = _BORDE

    for fila in ws.iter_rows(min_row=fila_encabezado + 1):
        for cell in fila:
            if cell.column in columnas_monto:
                cell.number_format = _FORMATO_BS

    for indice in range(1, len(columnas) + 1):
        largo = max(
            (len(str(fila[indice - 1].value or '')) for fila in ws.iter_rows(min_row=fila_encabezado)),
            default=0,
        )
        ws.column_dimensions[get_column_letter(indice)].width = min(largo + 4, 50)


def _respuesta_excel(wb, nombre_archivo):
    salida = io.BytesIO()
    wb.save(salida)
    respuesta = HttpResponse(
        salida.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    return respuesta


@login_required
@es_admin_o_tesorero
def egresos_excel(request):
    egresos, filtros = _egresos_filtrados(request)

    titulo = 'REPORTE DE EGRESOS'
    if filtros['fecha_inicio'] or filtros['fecha_fin']:
        titulo += f" — desde {filtros['fecha_inicio'] or 'el inicio'} hasta {filtros['fecha_fin'] or 'hoy'}"

    filas = []
    for e in egresos:
        registrado_por = ''
        if e.registrado_por:
            registrado_por = e.registrado_por.get_full_name() or e.registrado_por.username
        filas.append([
            e.pk,
            e.fecha.strftime('%d/%m/%Y'),
            e.categoria.nombre,
            e.descripcion,
            e.pagado_a,
            e.nro_respaldo,
            float(e.monto),
            f'Anulado: {e.motivo_anulacion}' if e.anulado else 'Vigente',
            registrado_por,
        ])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Egresos'
    _hoja_con_tabla(
        ws, titulo,
        ['N.º', 'Fecha', 'Categoría', 'Descripción', 'Pagado a', 'N.º factura/recibo',
         'Monto Bs', 'Estado', 'Registrado por'],
        filas,
        fila_total=['', '', '', '', '', 'TOTAL (sin anulados)', float(_total_egresos(egresos)), '', ''],
        columnas_monto=(7,),
    )
    return _respuesta_excel(wb, f'egresos_{timezone.localdate():%Y%m%d}.xlsx')


# =============================================================
# BALANCE: INGRESOS vs EGRESOS
# =============================================================

def _por_mes(consulta, campo_fecha, campo_monto):
    """Devuelve {numero_de_mes: total} para la consulta dada."""
    filas = consulta.values(mes=ExtractMonth(campo_fecha)).annotate(total=Sum(campo_monto))
    return {fila['mes']: fila['total'] or CERO for fila in filas}


def _datos_balance(anio):
    """
    Compara, mes a mes, lo que entró con lo que se gastó en el año.
    Los ingresos son los mismos del reporte de Recaudación: pagos aprobados
    más afiliaciones. Los egresos anulados no cuentan.
    """
    pagos = _por_mes(
        Pago.objects.filter(estado='Aprobado', fecha_pago__year=anio),
        'fecha_pago', 'monto_pagado')
    afiliaciones = _por_mes(
        Afiliacion.objects.filter(fecha_pago__year=anio),
        'fecha_pago', 'monto')
    egresos_vigentes = Egreso.objects.filter(anulado=False, fecha__year=anio)
    egresos = _por_mes(egresos_vigentes, 'fecha', 'monto')

    meses = []
    acumulado = CERO
    totales = {'pagos': CERO, 'afiliaciones': CERO, 'ingresos': CERO, 'egresos': CERO}

    for numero in range(1, 13):
        fila = {
            'nombre': MESES_NOMBRE[f'{numero:02d}'],
            'pagos': pagos.get(numero, CERO),
            'afiliaciones': afiliaciones.get(numero, CERO),
            'egresos': egresos.get(numero, CERO),
        }
        fila['ingresos'] = fila['pagos'] + fila['afiliaciones']
        fila['saldo'] = fila['ingresos'] - fila['egresos']
        acumulado += fila['saldo']
        fila['acumulado'] = acumulado
        meses.append(fila)

        for clave in totales:
            totales[clave] += fila[clave]

    totales['saldo'] = totales['ingresos'] - totales['egresos']

    por_categoria = list(
        egresos_vigentes.values('categoria__nombre')
        .annotate(total=Sum('monto'))
        .order_by('-total')
    )

    return {'meses': meses, 'totales': totales, 'por_categoria': por_categoria}


def _anio_pedido(request):
    anio = request.GET.get('anio', '').strip()
    if anio.isdigit() and 2000 <= int(anio) <= 2100:
        return int(anio)
    return timezone.localdate().year


@login_required
@es_admin_o_tesorero
def reporte_balance(request):
    anio = _anio_pedido(request)
    return render(request, 'reportes/balance.html', {
        'anio': anio,
        **_datos_balance(anio),
        **_metadata_reporte(request),
    })


@login_required
@es_admin_o_tesorero
def reporte_balance_excel(request):
    anio = _anio_pedido(request)
    datos = _datos_balance(anio)
    totales = datos['totales']

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Balance'
    _hoja_con_tabla(
        ws, f'BALANCE DE INGRESOS Y EGRESOS — GESTIÓN {anio}',
        ['Mes', 'Pagos de agua Bs', 'Afiliaciones Bs', 'Total ingresos Bs',
         'Egresos Bs', 'Saldo del mes Bs', 'Saldo acumulado Bs'],
        [
            [m['nombre'], float(m['pagos']), float(m['afiliaciones']), float(m['ingresos']),
             float(m['egresos']), float(m['saldo']), float(m['acumulado'])]
            for m in datos['meses']
        ],
        fila_total=['TOTAL', float(totales['pagos']), float(totales['afiliaciones']),
                    float(totales['ingresos']), float(totales['egresos']),
                    float(totales['saldo']), float(totales['saldo'])],
        columnas_monto=(2, 3, 4, 5, 6, 7),
    )

    ws = wb.create_sheet('Egresos por categoría')
    _hoja_con_tabla(
        ws, f'EGRESOS POR CATEGORÍA — GESTIÓN {anio}',
        ['Categoría', 'Total Bs'],
        [[c['categoria__nombre'], float(c['total'])] for c in datos['por_categoria']],
        fila_total=['TOTAL', float(totales['egresos'])],
        columnas_monto=(2,),
    )

    return _respuesta_excel(wb, f'balance_{anio}.xlsx')
